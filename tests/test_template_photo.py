"""模板照片字段、预填与原子另存为测试。"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from PIL import Image as PillowImage

from app.models.table_dataset import Provenance, TableDataset, TableRow
from app.parsers.workbook_template_analyzer import WorkbookTemplateAnalyzer
from app.services.workbook_fill_service import USE_NEW_VALUE, WorkbookFillService
from app.template_engine.generator import PhotoInsertError, TemplateGenerator
from app.template_engine.schema import FieldSchema, SheetSchema, TemplateSchema


def photo_schema(default_rows: int = 1) -> TemplateSchema:
    return TemplateSchema(
        "照片名单",
        default_rows=default_rows,
        sheets=[SheetSchema("录入", [FieldSchema("照片", "photo"), FieldSchema("姓名", "name", standard_field="name")])],
    )


def create_image(path: Path, format: str = "PNG") -> Path:
    PillowImage.new("RGB", (320, 160), "blue").save(path, format=format)
    return path


def test_photo_field_is_valid_and_allocates_input_area(tmp_path):
    path = TemplateGenerator().generate(photo_schema(), tmp_path / "照片模板.xlsx")
    sheet = load_workbook(path)["录入"]
    assert sheet.column_dimensions["A"].width >= 18
    assert sheet.row_dimensions[2].height >= 80


def test_missing_photo_does_not_leave_final_workbook(tmp_path):
    output = tmp_path / "结果.xlsx"
    with pytest.raises(PhotoInsertError, match="第 2 行"):
        TemplateGenerator().generate(photo_schema(), output, prefill_rows=[{"photo": tmp_path / "missing.png"}])
    assert not output.exists()
    assert not list(tmp_path.glob("*.partial.*"))


def test_prefill_rows_expand_beyond_default_rows_and_escape_text(tmp_path):
    rows = [{"name": "=学生" + str(index)} for index in range(5)]
    path = TemplateGenerator().generate(photo_schema(default_rows=2), tmp_path / "名单.xlsx", prefill_rows=rows)
    sheet = load_workbook(path)["录入"]
    assert sheet.max_row >= 6
    assert sheet["B2"].value == "'=学生0"


@pytest.mark.parametrize("suffix, format", [(".png", "PNG"), (".jpg", "JPEG")])
def test_prefilled_photo_is_reopenable_and_limited_to_cell_area(tmp_path, suffix, format):
    image = create_image(tmp_path / f"头像{suffix}", format)
    path = TemplateGenerator().generate(photo_schema(), tmp_path / f"照片{suffix}.xlsx", prefill_rows=[{"photo": image}])
    sheet = load_workbook(path)["录入"]
    assert len(sheet._images) == 1
    drawing = sheet._images[0].anchor.ext
    assert drawing.width / 9525 <= sheet.column_dimensions["A"].width * 7
    assert drawing.height / 9525 <= sheet.row_dimensions[2].height * 96 / 72


def test_compatibility_fill_inserts_limited_photo(tmp_path):
    image = create_image(tmp_path / "填充头像.png")
    template = tmp_path / "填写模板.xlsx"
    workbook = Workbook(); sheet = workbook.active; sheet.title = "名单"; sheet.append(["照片"]); workbook.save(template)
    dataset = TableDataset(["photo"], [TableRow({"photo": str(image)}, Provenance("虚构.xlsx", "名单", 2))], "虚构.xlsx", "名单", 1)
    analysis = WorkbookTemplateAnalyzer().analyze(template, "名单", 1)
    output = tmp_path / "填写结果.xlsx"

    result = WorkbookFillService().fill(analysis, dataset, {"照片": "photo"}, USE_NEW_VALUE, output_path=output, compatibility_accepted=True)

    sheet = load_workbook(result.output_path)["名单"]
    assert len(sheet._images) == 1
