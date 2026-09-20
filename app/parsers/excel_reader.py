"""xlsx/csv 只读加载器，不修改源文件。"""
from dataclasses import dataclass, field
from datetime import datetime
from numbers import Number
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles.numbers import is_date_format


@dataclass(frozen=True)
class ExcelCellDateContext:
    """单元格已由源工作簿明确标记为 Excel 日期的上下文。"""

    date_semantic: bool
    excel_epoch: datetime | None


@dataclass(frozen=True)
class RawWorkbookData:
    """原始表格与仅在读取阶段提取的日期语义。"""

    frame: pd.DataFrame
    date_context: dict[tuple[int, int], ExcelCellDateContext] = field(default_factory=dict)

SUPPORTED={".xlsx",".csv"}
def list_sheets(path:Path)->list[str]:
    if path.suffix.lower()==".csv": return []
    if path.suffix.lower()!=".xlsx": raise ValueError("仅支持 .xlsx 和 .csv 文件")
    try: return pd.ExcelFile(path,engine="openpyxl").sheet_names
    except Exception as e: raise ValueError(f"无法读取 Excel 文件：{e}") from e
def read_raw(path:Path,sheet_name:str|None=None)->pd.DataFrame:
    if not path.exists(): raise ValueError("文件不存在")
    ext=path.suffix.lower()
    if ext==".xlsx":
        try: frame=pd.read_excel(path,sheet_name=sheet_name or 0,header=None,dtype=object,engine="openpyxl")
        except Exception as e: raise ValueError(f"无法读取 Excel 文件：{e}") from e
    elif ext==".csv":
        for encoding in ("utf-8-sig","gbk","utf-8"):
            try: frame=pd.read_csv(path,header=None,dtype=object,encoding=encoding); break
            except UnicodeDecodeError: continue
        else: raise ValueError("CSV 编码无法识别")
    else: raise ValueError("仅支持 .xlsx 和 .csv 文件")
    if frame.empty: raise ValueError("文件或工作表为空")
    return frame


def read_raw_with_context(path: Path, sheet_name: str | None = None) -> RawWorkbookData:
    """读取原始数据，并在 `.xlsx` 的读取阶段保留日期单元格语义。"""
    frame = read_raw(path, sheet_name)
    if path.suffix.lower() != ".xlsx":
        return RawWorkbookData(frame)

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            worksheet = workbook[sheet_name or workbook.sheetnames[0]]
            context: dict[tuple[int, int], ExcelCellDateContext] = {}
            for row_index in range(frame.shape[0]):
                for column_index in range(frame.shape[1]):
                    cell = worksheet.cell(row=row_index + 1, column=column_index + 1)
                    if cell.data_type not in {"n", "d"} or not is_date_format(cell.number_format):
                        continue
                    value = frame.iat[row_index, column_index]
                    is_numeric = isinstance(value, Number) and not isinstance(value, bool)
                    if is_numeric or cell.data_type == "d":
                        context[(row_index, column_index)] = ExcelCellDateContext(True, workbook.epoch)
            return RawWorkbookData(frame, context)
        finally:
            workbook.close()
    except Exception as error:
        raise ValueError(f"无法读取 Excel 日期语义：{error}") from error
